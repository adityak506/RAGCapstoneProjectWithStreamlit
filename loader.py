# utils/loader.py
from langchain_community.document_loaders import TextLoader, PyPDFLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from docx import Document as DocxDocument
import os
import config
import llm_client

def load_and_chunk_docs(folder_path = 'DATA_DIR', chunk_size=900, chunk_overlap=150):
    docs = []

    for file in os.listdir(folder_path):
        path = os.path.join(folder_path, file)

        if file.endswith(".txt"):
            loader = TextLoader(path)
        elif file.endswith(".docx"):
            docx = DocxDocument(path)
            docs.append(
                Document(
                    page_content="\n".join(paragraph.text for paragraph in docx.paragraphs),
                    metadata={"source": path},
                )
            )
            continue
        elif file.endswith(".pdf"):
            loader = PyPDFLoader(path)
        else:
            continue

        docs.extend(loader.load())

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap
    )
    chunks = splitter.split_documents(docs)

    print(f"✅ Loaded {len(docs)} docs → {len(chunks)} chunks")
    return chunks


if __name__ == "__main__":
    load_and_chunk_docs("data")