"""Generate the fictional DOCX and PDF documents included in this project."""

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


DATA_DIR = Path(__file__).resolve().parent / "data"

AI_GUIDE_SECTIONS = {
    "Purpose and Team Structure": [
        "NovaTech's AI Engineering Practice builds reliable machine-learning and "
        "generative-AI features for OrbitOps, InsightFlow, and customer projects. "
        "Each delivery team usually includes an AI engineer, data engineer, product "
        "owner, domain reviewer, and security champion.",
        "Teams use short, reviewable experiments. A prototype must have a written "
        "problem statement, success metric, data owner, and named reviewer before it "
        "can move into product development.",
    ],
    "The AI Project Lifecycle": [
        "Discovery begins with a one-page AI opportunity brief. The team records the "
        "user problem, expected value, known risks, and a non-AI baseline. During data "
        "readiness, engineers document data origin, consent, retention, quality, and "
        "known gaps. Experiments are tracked with model version, prompt version, dataset "
        "version, parameters, and evaluation results.",
        "A project can enter a pilot only after the product owner accepts the usefulness "
        "evaluation and the security champion accepts the risk review. Production release "
        "also requires monitoring, rollback instructions, an incident owner, and a model "
        "or prompt card. Teams review production AI systems at least once each quarter.",
    ],
    "Retrieval-Augmented Generation Standard": [
        "NovaTech uses retrieval-augmented generation when answers must be grounded in "
        "approved company or customer material. The basic workflow is load, clean, chunk, "
        "embed, retrieve, build context, and generate. Teams must preserve source metadata "
        "so users can inspect the evidence behind an answer.",
        "The starting chunk size is 600 to 1,000 characters with 10 to 20 percent overlap. "
        "Engineers tune these values using representative questions rather than assuming "
        "one setting works for every collection. Retrieval evaluations track context "
        "precision, context recall, answer correctness, citation correctness, and the rate "
        "of appropriate 'not found' responses.",
        "For conversational search, a short query-rewriting prompt may convert follow-up "
        "questions into standalone retrieval queries. Retrieved chunks should be filtered "
        "by a tested relevance threshold before they reach the answer model. Threshold "
        "values are embedding-model-specific and must be re-evaluated after a model change.",
    ],
    "Prompting and Evaluation": [
        "Prompts should state the task, permitted context, output format, and refusal "
        "behavior in plain language. Prompts are versioned with the application code. "
        "A prompt change requires a small regression set covering typical, difficult, "
        "ambiguous, and unanswerable questions.",
        "Human review is required for evaluations involving safety, legal interpretation, "
        "or high-impact customer decisions. Synthetic test data is encouraged, but it must "
        "not replace representative, authorized examples. Evaluation results are attached "
        "to the release record.",
    ],
    "Security Requirements for AI Teams": [
        "AI engineers may use only NovaTech-approved model providers and accounts. Customer "
        "data, source code, credentials, personal data, and internal restricted documents "
        "must never be pasted into public consumer AI tools. API keys belong in the approved "
        "secret manager or local environment variables, never in notebooks or Git.",
        "Before sending data to a model API, teams minimize the fields, remove direct "
        "identifiers where possible, and confirm the provider's retention setting. RAG "
        "indexes inherit the highest classification of their source documents. Access to "
        "an index must match access to its sources.",
        "Prompt-injection testing is mandatory for RAG applications that read user-supplied "
        "documents. Retrieved text is treated as untrusted data, not as system instructions. "
        "Applications must log document identifiers and model versions without logging "
        "sensitive prompt content.",
    ],
    "Remote Collaboration for AI Engineering": [
        "AI engineering teams follow the company hybrid-work policy in the employee policy "
        "document. In addition, remote experiments must run on managed NovaTech devices or "
        "approved cloud workspaces. Restricted datasets may not be downloaded to personal "
        "devices, removable media, or unapproved local folders.",
        "Teams hold a 15-minute weekday checkpoint in their agreed collaboration window. "
        "Design decisions, experiment results, and handoffs must be written in the project "
        "workspace so colleagues in other time zones are not excluded. Production model "
        "changes cannot be approved solely through a chat message.",
    ],
    "Incident Response and Ownership": [
        "Examples of AI incidents include disclosure of sensitive context, harmful or "
        "materially incorrect output, unexplained model drift, broken citations, or sudden "
        "cost spikes. The on-call engineer disables the affected feature or switches to "
        "the documented fallback, preserves relevant evidence, and opens a security "
        "incident ticket.",
        "The team notifies the Security Response Desk within 30 minutes for a suspected "
        "data exposure and within four hours for other serious AI incidents. A blameless "
        "review documents impact, cause, correction, and prevention actions.",
    ],
}

POLICY_PAGES = [
    (
        "1. Employment, Conduct, and Leave",
        [
            "NovaTech Solutions provides 20 days of paid annual leave each calendar year. "
            "Employees request planned leave in PeopleHub at least five business days in "
            "advance when practical. Managers normally respond within two business days. "
            "Unused annual leave may carry over up to five days and must be used by March 31.",
            "Employees receive 10 days of paid sick leave. For absences longer than three "
            "consecutive workdays, People Operations may request appropriate documentation. "
            "NovaTech also provides 12 weeks of paid parental leave after six months of "
            "employment and three paid volunteer days per year.",
            "All employees must act respectfully, protect confidential information, avoid "
            "conflicts of interest, and report suspected misconduct through a manager, "
            "People Operations, or the confidential SpeakUp channel.",
        ],
    ),
    (
        "2. Hybrid and Remote Work Policy",
        [
            "Roles marked hybrid normally require two anchor days per week at the employee's "
            "assigned hub. Teams publish their anchor days each quarter. Fully remote roles "
            "are allowed when the employment agreement explicitly says remote or when People "
            "Operations approves an accommodation.",
            "Employees working remotely must be reachable during the team's four-hour core "
            "collaboration window. They must use a company-managed device, connect through "
            "the NovaSecure VPN on untrusted networks, keep screens private, and avoid "
            "confidential conversations in public spaces.",
            "Temporary work from another domestic location is permitted for up to 20 working "
            "days per calendar year with manager approval. International remote work requires "
            "People Operations, tax, and security approval before travel; it is not automatic.",
            "AI engineering teams also follow the stricter data and experiment controls in "
            "the NovaTech AI Engineering Guide. If two policies differ, the more protective "
            "security rule applies.",
        ],
    ),
    (
        "3. Information Classification and Access",
        [
            "NovaTech information has four levels: Public, Internal, Confidential, and "
            "Restricted. Public content is approved for anyone. Internal content is for "
            "NovaTech workers. Confidential content includes customer contracts, non-public "
            "roadmaps, and employee records. Restricted content includes credentials, "
            "production secrets, regulated personal data, and customer encryption keys.",
            "Access follows least privilege and is reviewed every quarter for production "
            "systems. Managers approve business access; system owners approve technical "
            "access. Restricted access also requires security approval and multi-factor "
            "authentication. Access must be removed within four hours after an employee's "
            "departure is recorded.",
            "Confidential and Restricted information must be encrypted in transit and at "
            "rest. Restricted information may be stored only in approved systems and must "
            "not be copied into personal email, consumer storage, or unapproved AI tools.",
        ],
    ),
    (
        "4. Security Baseline",
        [
            "Multi-factor authentication is mandatory for email, source control, cloud "
            "systems, VPN, and administrative tools. Passwords must be unique and stored in "
            "the approved password manager. Sharing accounts or authentication factors is "
            "prohibited.",
            "Company devices use full-disk encryption, endpoint protection, automatic screen "
            "lock after ten minutes, and centrally managed security updates. Critical patches "
            "must be applied within seven days; high-severity patches within 14 days.",
            "Employees report suspected phishing, lost devices, exposed credentials, or "
            "unusual account activity immediately using the Security button in HelpDesk or "
            "by calling the Security Response Desk. A lost device should be reported within "
            "one hour. Employees should not investigate an incident on their own.",
        ],
    ),
    (
        "5. Customer Data and Retention",
        [
            "Customer data is used only for the documented service purpose and according to "
            "the customer agreement. Production customer data must not be used in development "
            "or AI evaluation unless the data owner and Security approve a documented exception.",
            "Support attachments are retained for 90 days after ticket closure unless a "
            "contract requires a different period. Security audit logs are retained for "
            "13 months. Project teams define retention for derived datasets and vector indexes "
            "before launch; deletion of source data must trigger deletion or rebuilding of "
            "derived indexes where applicable.",
            "A suspected customer-data exposure is a priority-one incident. Notify the Security "
            "Response Desk immediately, preserve evidence, and do not contact the customer "
            "unless the incident commander authorizes communication.",
        ],
    ),
]


def create_docx(output_path: Path) -> None:
    """Create the multi-section NovaTech AI Engineering Guide.

    Args:
        output_path: Destination path for the DOCX file.

    Returns:
        Nothing. A Word document is written to disk.
    """
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    document.add_heading("NovaTech AI Engineering Guide", 0)
    document.add_paragraph(
        "Version 2.1 · Owner: AI Engineering Practice · Fictional training material"
    )
    for heading, paragraphs in AI_GUIDE_SECTIONS.items():
        document.add_heading(heading, level=1)
        for text in paragraphs:
            paragraph = document.add_paragraph(text)
            paragraph.paragraph_format.space_after = Pt(8)
    document.save(output_path)


def create_pdf(output_path: Path) -> None:
    """Create a five-page employee and security policy PDF.

    Args:
        output_path: Destination path for the PDF file.

    Returns:
        Nothing. A searchable text PDF is written to disk.
    """
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CenteredTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        spaceAfter=18,
    )
    body_style = ParagraphStyle(
        "PolicyBody",
        parent=styles["BodyText"],
        fontSize=10.5,
        leading=15,
        spaceAfter=12,
    )
    pdf = SimpleDocTemplate(
        str(output_path),
        pagesize=LETTER,
        rightMargin=0.8 * inch,
        leftMargin=0.8 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
        title="NovaTech Employee and Security Policies",
        author="NovaTech Solutions (fictional)",
    )
    story = []
    for page_number, (heading, paragraphs) in enumerate(POLICY_PAGES):
        story.append(Paragraph("NovaTech Employee &amp; Security Policies", title_style))
        story.append(Paragraph(heading, styles["Heading1"]))
        story.append(
            Paragraph(
                "Effective January 2026 · Fictional material for classroom use",
                styles["Italic"],
            )
        )
        story.append(Spacer(1, 14))
        for text in paragraphs:
            story.append(Paragraph(text, body_style))
        if page_number < len(POLICY_PAGES) - 1:
            story.append(PageBreak())
    pdf.build(story)


def main() -> None:
    """Generate both binary sample documents in the data directory.

    Returns:
        Nothing. The function prints the generated file paths.
    """
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    docx_path = DATA_DIR / "ai_engineering_guide.docx"
    pdf_path = DATA_DIR / "employee_policies.pdf"
    create_docx(docx_path)
    create_pdf(pdf_path)
    print(f"Created {docx_path}")
    print(f"Created {pdf_path}")


if __name__ == "__main__":
    main()
