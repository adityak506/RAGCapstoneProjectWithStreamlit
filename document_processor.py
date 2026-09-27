"""Load, clean, and chunk the TXT, DOCX, and PDF knowledge-base files."""

import re
from pathlib import Path
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".txt", ".docx", ".pdf"}

def discover_documents(data_dir: Path) -> list[Path]:
    """Find supported knowledge-base files in a directory.

    Args:
        data_dir: Folder that contains the source documents.

    Returns:
        A sorted list of TXT, DOCX, and PDF file paths.
    """
    if not data_dir.exists():
        raise FileNotFoundError(f"Data directory does not exist: {data_dir}")
    return sorted(
        path
        for path in data_dir.iterdir()
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def clean_text(text: str) -> str:
    """Remove distracting whitespace while preserving paragraph breaks.

    Cleaning gives the embedding model readable text without flattening every
    paragraph into one long line.

    Args:
        text: Raw text extracted from a source document.

    Returns:
        Text with normalized spaces and blank lines.
    """
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    # Replace null characters and Windows/Mac line endings with normal newline characters for consistent text cleanup.

    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    # Split text into lines, collapse repeated spaces/tabs in each line, and trim leading/trailing whitespace.

    cleaned = "\n".join(lines)
    # Rejoin the cleaned lines with newline separators to preserve paragraph structure.

    return re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    # Remove excessive blank lines and trim leading/trailing whitespace from the final cleaned text.


def load_text_file(file_path: Path) -> list[Document]:
    """Load one UTF-8 text file as a LangChain document.

    Args:
        file_path: Path to the text file that should be read.

    Returns:
        A one-item list containing its text and source metadata.
    """
    text = clean_text(file_path.read_text(encoding="utf-8"))
    if not text:
        return []
    return [
        Document(
            page_content=text,
            metadata={"source": file_path.name, "file_type": "txt"},
        )
    ]


def load_docx_file(file_path: Path) -> list[Document]:
    """Load a DOCX file and retain its headings as simple section metadata.

    Args:
        file_path: Path to the Word document that should be read.

    Returns:
        One LangChain document per non-empty heading section.
    """
    docx = DocxDocument(file_path)
    documents: list[Document] = []
    section_name = "Introduction"
    section_paragraphs: list[str] = []

    def save_section() -> None:
        """Save the current DOCX section when it contains readable text."""
        text = clean_text("\n\n".join(section_paragraphs))
        if text:
            documents.append(
                Document(
                    page_content=f"{section_name}\n\n{text}", #page content has introduction then the paragraph
                    metadata={
                        "source": file_path.name,
                        "file_type": "docx",
                        "section": section_name,
                    },
                )
            )

    for paragraph in docx.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.startswith("Heading"): #since we want to call each and every section i.e. why Heading
            save_section()
            section_name = text
            section_paragraphs = []
        else:
            section_paragraphs.append(text)
    save_section()
    return documents


def load_pdf_file(file_path: Path) -> list[Document]:
    """Load a PDF as one LangChain document per non-empty page.

    Args:
        file_path: Path to the PDF document that should be read.

    Returns:
        Documents containing page text plus source and page-number metadata.
    """
    reader = PdfReader(str(file_path))
    documents: list[Document] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = clean_text(page.extract_text() or "")
        if text:
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": file_path.name,
                        "file_type": "pdf",
                        "page": page_number,
                    },
                )
            )
    return documents


def load_documents(data_dir: Path) -> list[Document]:
    """Load every supported document in the knowledge-base directory.

    Args:
        data_dir: Folder containing TXT, DOCX, and PDF files.

    Returns:
        Loaded documents with metadata. Empty files are skipped.
    """
    loaders = {
        ".txt": load_text_file,
        ".docx": load_docx_file,
        ".pdf": load_pdf_file,
    }
    file_paths = discover_documents(data_dir) 
    # discover_documents TO GET THE SORTED PATH OF THE ALLOWED or supported DOCUMENTS, function which we created above
    
    if not file_paths:
        raise FileNotFoundError(f"No supported documents were found in {data_dir}.")

    documents: list[Document] = []
    for file_path in file_paths:
        documents.extend(loaders[file_path.suffix.lower()](file_path)) #we give extension of the file path in the loaders dict.
#loaders is a dictionary where we give an extension and returns the function like load_text_file for loading files with dif. ext.
    if not documents:
        raise ValueError("The source files were found, but no readable text was extracted.")
    print('documents: ', documents)
    return documents


def chunk_documents(
    documents: list[Document],
    chunk_size: int,
    chunk_overlap: int,
) -> list[Document]:
    """Split long documents into smaller, overlapping pieces for retrieval.

    Embeddings work best on focused passages. ``chunk_size`` limits each
    passage, while ``chunk_overlap`` repeats a little text so ideas near a
    boundary are not lost.

    Args:
        documents: Loaded documents that need to be split.
        chunk_size: Maximum approximate number of characters per chunk.
        chunk_overlap: Number of characters shared by neighboring chunks.

    Returns:
        Smaller documents with original metadata and a unique chunk number.
    """
    if chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap must be smaller than chunk_size.")
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(documents)
    for chunk_number, chunk in enumerate(chunks, start=1):
        chunk.metadata["chunk_id"] = chunk_number
    return chunks